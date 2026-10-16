#!/usr/bin/env python3
"""Создаёт НОВЫЙ Excel-файл запуска с номером и датой: NNN_resellers_ГГГГ-ММ-ДД.xlsx.

Использование:
  python3 build_xlsx.py rows.txt /путь/к/папке [--run-note "Camry/Rio/Solaris, moskovskaya_oblast"]

rows.txt - строки из Приложения C (по одной на продавца, поля через |).
Папка    - папка на рабочем столе (её создаёт setup_folder.py). Номер NNN = следующий по порядку (001, 002, ...),
           дата = сегодняшняя; два запуска в один день дают два файла с разными номерами.
Дубли отсекаются по ID профиля среди ВСЕХ файлов папки (старых запусков тоже). Если новых нет, файл не создаётся.
Печатает JSON: file, added, duplicates_skipped. Телефонов в таблице нет и быть не должно.
"""
import sys, json, re, datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

F = 'Arial'
HDR = ['№', 'Дата добавления', 'Запрос (модель)', 'Продавец (ник)', 'ID профиля', 'Ссылка на профиль',
       'Ссылка на объявление', 'Гараж: всего', 'Гараж: без «бывших»', 'Что в гараже', 'Уверенность',
       'Дубль?', 'Статус контакта', 'Дата контакта', 'Комментарий']
STATUSES = '"не писали,написали,ответил,отказ,не отвечает,не перекуп"'
FILE_RE = re.compile(r'^(\d{3})_resellers_(\d{4}-\d{2}-\d{2})\.xlsx$')
PROFILE_RE = re.compile(r'auto\.ru/profile/([^/\s?#]+)')

def parse(path):
    out = []
    for line in Path(path).read_text(encoding='utf-8').splitlines():
        p = line.split('|')
        if len(p) < 7:
            continue
        q, pid, name, gt, gc, g, l = p[:7]
        out.append(dict(q=q, pid=pid.strip(), name=name, gt=int(gt), gc=int(gc), g=g, l=l.strip()))
    return out

def confidence(gc):
    # гараж не равен объявлениям на продаже, поэтому честно называем это оценкой
    return 'средняя (3+ машин)' if gc >= 3 else 'низкая-средняя (2 машины, проверить)'

def existing_ids(folder):
    """ID профилей из всех .xlsx папки (старый resellers.xlsx и файлы с номерами тоже)."""
    ids = set()
    for f in folder.glob('*.xlsx'):
        if f.name.startswith('~$'):
            continue
        try:
            ws = load_workbook(f)['Продавцы']
        except Exception:
            continue
        hdr = {str(c.value): c.column for c in ws[1] if c.value}
        col = hdr.get('ID профиля')
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if col and c.column == col and c.value:
                    ids.add(str(c.value).strip())
                elif not col:
                    s = (c.hyperlink.target if c.hyperlink and c.hyperlink.target else c.value)
                    m = PROFILE_RE.search(s) if isinstance(s, str) else None
                    if m:
                        ids.add(m.group(1))
    return ids

def main():
    rows_path, folder = sys.argv[1], Path(sys.argv[2]).expanduser()
    note = sys.argv[sys.argv.index('--run-note') + 1] if '--run-note' in sys.argv else ''
    folder.mkdir(parents=True, exist_ok=True)
    rows = parse(rows_path)
    seen = existing_ids(folder)
    today = datetime.date.today().isoformat()
    fresh, skipped = [], 0
    for x in rows:
        if x['pid'] in seen:
            skipped += 1; continue
        seen.add(x['pid']); fresh.append(x)
    if not fresh:
        print(json.dumps({'file': None, 'added': 0, 'duplicates_skipped': skipped}, ensure_ascii=False)); return
    nums = [int(m.group(1)) for n in (p.name for p in folder.iterdir()) for m in [FILE_RE.match(n)] if m]
    path = folder / f'{(max(nums) + 1 if nums else 1):03d}_resellers_{today}.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Продавцы'; ws.append(HDR)
    for i, x in enumerate(fresh, 1):
        prof = f"https://auto.ru/profile/{x['pid']}/"
        lst = 'https://auto.ru/cars/used/sale/' + x['l'].lstrip('/')
        ws.append([i, today, x['q'], x['name'], x['pid'], prof, lst, x['gt'], x['gc'], x['g'],
                   confidence(x['gc']), None, 'не писали', '', ''])
        ws.cell(i + 1, 6).hyperlink = prof; ws.cell(i + 1, 7).hyperlink = lst
    last = ws.max_row
    for r in range(2, last + 1):
        ws.cell(r, 12).value = f'=IF(COUNTIF($E$2:$E${last},E{r})>1,"ДУБЛЬ","ок")'
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.font = Font(name=F, color='0563C1', underline='single') if c.column in (6, 7) else Font(name=F)
            c.alignment = Alignment(vertical='top', wrap_text=c.column in (10, 15))
    for col, w in zip('ABCDEFGHIJKLMNO', [5, 13, 18, 22, 26, 30, 36, 10, 12, 44, 22, 9, 16, 13, 28]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:O{last}'
    dv = DataValidation(type='list', formula1=STATUSES, allow_blank=True); ws.add_data_validation(dv); dv.add('M2:M2000')
    log = wb.create_sheet('Журнал запусков')
    log.append(['Дата', 'Критерии и запрос', 'Найдено', 'Добавлено', 'Дублей пропущено'])
    log.append([today, note, len(rows), len(fresh), skipped])
    for c in log[1]:
        c.font = Font(name=F, bold=True)
    log.column_dimensions['B'].width = 70
    wb.save(path)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps({'file': str(path), 'added': len(fresh), 'duplicates_skipped': skipped}, ensure_ascii=False))

if __name__ == '__main__':
    main()
