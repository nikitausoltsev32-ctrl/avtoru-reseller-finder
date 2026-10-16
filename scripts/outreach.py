#!/usr/bin/env python3
"""Помощник рассылки: выбирает, кому писать, и ведёт отметки в Excel-файлах папки AutoRu_Resellers.
Сам ничего не отправляет: отправка идёт в браузере, только после подтверждения пользователя в чате.

Команды:
  python3 outreach.py pick  <папка> --limit N [--daily-cap 10] [--only-strong]
      Печатает JSON: сколько уже отправлено сегодня, сколько можно ещё, и список кандидатов
      (статус «не писали», без «снято», сначала «средняя» уверенность, затем старые файлы).
  python3 outreach.py mark  <папка> --pid ID --status написали [--comment "..."]
      Ставит статус и дату контакта в нужном файле. Допустимые статусы: написали, ответил, отказ,
      не отвечает, не перекуп.
  python3 outreach.py skip  <папка> --pid ID --comment "объявление снято"
      Оставляет «не писали», но помечает строку, чтобы не выбирать её снова.
  python3 outreach.py message <папка> [--set "текст"]
      Показывает (или сохраняет) текст сообщения из message.txt.
"""
import sys, json, re, datetime
from pathlib import Path
from openpyxl import load_workbook

STATUSES = {'написали', 'ответил', 'отказ', 'не отвечает', 'не перекуп'}
HARD_CAP = 20          # больше этого в сутки скрипт не выдаёт никогда
COL = dict(n=1, date=2, model=3, nick=4, pid=5, prof=6, listing=7, conf=11, status=13, cdate=14, comment=15)
FILE_RE = re.compile(r'^(\d{3})_resellers_\d{4}-\d{2}-\d{2}\.xlsx$')

def files(folder):
    return sorted(p for p in Path(folder).expanduser().glob('*.xlsx')
                  if not p.name.startswith('~$') and (FILE_RE.match(p.name) or p.name == 'resellers.xlsx'))

def arg(name, default=None):
    return sys.argv[sys.argv.index(name) + 1] if name in sys.argv else default

def rows(folder):
    for f in files(folder):
        try:
            ws = load_workbook(f)['Продавцы']
        except Exception:
            continue
        for r in range(2, ws.max_row + 1):
            pid = ws.cell(r, COL['pid']).value
            if pid:
                yield f, r, ws, {k: ws.cell(r, c).value for k, c in COL.items()}

def cmd_pick(folder):
    today = datetime.date.today().isoformat()
    limit = int(arg('--limit', '5'))
    cap = min(int(arg('--daily-cap', '10')), HARD_CAP)
    sent_today, cands = 0, []
    for f, r, ws, d in rows(folder):
        if str(d['cdate'] or '')[:10] == today and d['status'] in ('написали', 'ответил', 'отказ', 'не отвечает'):
            sent_today += 1
        status = (d['status'] or 'не писали').strip()
        if status != 'не писали' or 'снято' in str(d['comment'] or '').lower():
            continue
        strong = str(d['conf'] or '').startswith('средняя')
        if '--only-strong' in sys.argv and not strong:
            continue
        cands.append((0 if strong else 1, f.name, r, dict(
            pid=str(d['pid']).strip(), nick=d['nick'], model=d['model'], confidence=d['conf'],
            listing=d['listing'], profile=d['prof'], file=f.name)))
    left = max(0, cap - sent_today)
    take = [c[3] for c in sorted(cands, key=lambda c: (c[0], c[1], c[2]))[:min(limit, left)]]
    print(json.dumps(dict(sent_today=sent_today, daily_cap=cap, can_send_more=left,
                          candidates_total=len(cands), picked=take), ensure_ascii=False))

def update(folder, pid, status=None, comment=None):
    for f, r, ws, d in rows(folder):
        if str(d['pid']).strip() == pid:
            wb = load_workbook(f); w = wb['Продавцы']
            if status:
                w.cell(r, COL['status']).value = status
                w.cell(r, COL['cdate']).value = datetime.date.today().isoformat()
            if comment:
                old = w.cell(r, COL['comment']).value
                w.cell(r, COL['comment']).value = (str(old) + '; ' if old else '') + comment
            wb.save(f)
            print(json.dumps(dict(ok=True, file=f.name, row=r, pid=pid, status=status), ensure_ascii=False)); return
    print(json.dumps(dict(ok=False, error='pid_not_found', pid=pid), ensure_ascii=False)); sys.exit(1)

def cmd_message(folder):
    p = Path(folder).expanduser() / 'message.txt'
    text = arg('--set')
    if text is not None:
        p.write_text(text, encoding='utf-8')
    print(json.dumps(dict(path=str(p), text=p.read_text(encoding='utf-8') if p.exists() else None), ensure_ascii=False))

if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    cmd, folder = sys.argv[1], sys.argv[2]
    if cmd == 'pick': cmd_pick(folder)
    elif cmd == 'mark':
        st = arg('--status')
        if st not in STATUSES: sys.exit('status must be one of: ' + ', '.join(sorted(STATUSES)))
        update(folder, arg('--pid'), st, arg('--comment'))
    elif cmd == 'skip': update(folder, arg('--pid'), None, arg('--comment', 'объявление снято'))
    elif cmd == 'message': cmd_message(folder)
    else: sys.exit('unknown command')
