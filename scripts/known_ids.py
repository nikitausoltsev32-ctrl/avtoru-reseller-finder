#!/usr/bin/env python3
"""Печатает ID профилей из ВСЕХ Excel-файлов папки одной строкой JSON-массива:
вставьте её в CONFIG.knownProfileIds сборщика (Приложение A).
Использование: python3 known_ids.py /путь/к/папке   (можно передать и один .xlsx)
Читает все .xlsx в папке (кроме временных ~$...), ищет ссылки/ID вида auto.ru/profile/<ID> на листе «Продавцы»."""
import sys, json, re
from pathlib import Path
from openpyxl import load_workbook

PROFILE_RE = re.compile(r'auto\.ru/profile/([^/\s?#]+)')

def known_ids(target):
    t = Path(target)
    files = [t] if t.is_file() else sorted(p for p in t.glob('*.xlsx') if not p.name.startswith('~$'))
    ids = []
    for f in files:
        try:
            wb = load_workbook(f)
            ws = wb['Продавцы']
        except Exception:
            continue
        hdr = {str(c.value): c.column for c in ws[1] if c.value}
        col = hdr.get('ID профиля')
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if col and c.column == col and c.value:
                    ids.append(str(c.value).strip())
                elif c.hyperlink and c.hyperlink.target:
                    m = PROFILE_RE.search(c.hyperlink.target)
                    if m and not col:
                        ids.append(m.group(1))
                elif isinstance(c.value, str) and not col:
                    m = PROFILE_RE.search(c.value)
                    if m:
                        ids.append(m.group(1))
    return list(dict.fromkeys(ids))

if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(known_ids(sys.argv[1]) if len(sys.argv) > 1 else [], ensure_ascii=False))
